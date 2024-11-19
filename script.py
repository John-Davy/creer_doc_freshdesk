import os
import logging
import subprocess
import json
from docx import Document
from datetime import datetime
from api import add_attachment_to_ticket, resquest_product_name, resquest_collaborateur_info, add_note_to_ticket
from collections import defaultdict

# Définir le chemin du template
TEMPLATE = defaultdict(lambda: "Le type d'actif reçu n'est pas pris en charge",{
    "Mobile Phone" : './templates/template_recommandations_Mobile_Phone.docx',
    "Laptop" : './templates/template_installer_Laptop_ou_Tablet.docx',
    "Tablet" : './templates/template_installer_Laptop_ou_Tablet.docx'
})

TICKET_ID = ''
            
# Initialisation de l'indicateur d'erreur
ERRORS_OCCURED = False      

# Renvoi un nom de fichier unique
def get_unique_filename(file_path):
    """ Génère un nom de fichier unique en ajoutant un numéro incrémental si nécessaire. """
    base, ext = os.path.splitext(file_path)
    counter = 1
    new_file_path = file_path
    while os.path.isfile(new_file_path):
        new_file_path = f"{base}_{counter}{ext}"
        counter += 1
    return new_file_path
    
def format_date(date_str):
    """ Convertit une date en chaîne au format 'DD.MM.YYYY' ou retourne une valeur par défaut """
    try:
        date_obj = datetime.strptime(date_str, '%a, %d %b, %Y %H:%M GMT %z')
        return date_obj.strftime('%d.%m.%Y')
    except ValueError as e:
        content_note = (f"script.py :\nformat_date: Erreur de formatage de la date : {e}") 
        logging.error(content_note, exc_info=True)
        add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        ERRORS_OCCURED = True
        return str(datetime.now().strftime('%d.%m.%Y'))

def open_doc(doc_path):
    """ Crée une instance du template et la renvoie """
    try:
        # Vérifie si le chemin du template existe
        if not os.path.isfile(doc_path):
            content_note = f"script.py :\nopen_doc : Le fichier template n'existe pas à l'emplacement : {doc_path}"
            logging.error(content_note)
            add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
            raise FileNotFoundError("script.py - open_doc")
        # Crée une instance doc avec le template situé à doc_path
        doc = Document(doc_path)
        return doc
    except Exception as e:
        content_note = f"script.py :\nopen_doc : Erreur lors du chargement du document : {e}"
        logging.error(content_note, exc_info=True)
        add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        ERRORS_OCCURED = True
        raise FileNotFoundError("script.py - open_doc")

def create_json(raw_data):
    """ Converti la chaine de caractère et renvoie un object JSON """
    try:
        data = json.loads(raw_data)
    except json.JSONDecodeError as e:
        content_note = f"script.py :\ncreate_json : Erreur lors du parsing du JSON : {e}"
        logging.error(content_note, exc_info=True)
        add_note_to_ticket(TICKET_ID, "create_json : Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        ERRORS_OCCURED = True
        raise Exception()
    except Exception as e:
        content_note = f"script.py :\ncreate_json : Erreur lors de l'obtention du JSON : {e}"
        logging.error(content_note, exc_info=True)
        add_note_to_ticket(TICKET_ID, "create_json : Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        ERRORS_OCCURED = True
        raise Exception()
    return data
    
def create_dir(path):
    """ Crée de répertoire si inexistant """
    try:
        if not os.path.exists(path):
            os.makedirs(path)
    except Exception as e:
        content_note = f"script.py :\ncreate_dir : Erreur lors de la création du répertoire {path} : {e}"
        logging.error(content_note, exc_info=True)
        add_note_to_ticket(TICKET_ID, "create_dir : Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        ERRORS_OCCURED = True
        raise Exception()
        
def init_palceholders(data, custom_fields):
    """ Reçoit les Json data et custom_fields et initialise le dictionnaire palceholders et le renvoi en sortie """
    placeholders = {}
    
    # gestion dela date
    date = str(data.get('Date'))
    date = format_date(date)
    placeholders["{{Date}}"] = date
    
    # valeurs obligatoires
    requiered_values = {
        "{{Utilise_par}}" : data.get("Used_by"),
        "{{Asset_tag}}" : data.get("Asset_tag"), 
        "{{Numéro_de_série}}" : custom_fields.get("serial_number_50000227369")
    }
    
    for key, value in requiered_values.items():
        if value not in [ None, ""] :
            placeholders[key] = value
        else :
            content_note = f"La variable {key} n'a pas de valeur !"
            logging.error(f"script.py :\ninit_palceholders : {content_note}")
            ERRORS_OCCURED = True
            add_note_to_ticket(TICKET_ID, content_note)
    return placeholders

def create_placeholders(cl_type, data, custom_fields):
    """ Crée le dictionnaire avec les bonnes clefs et valeurs pour remplir le doc """
    placeholders = init_palceholders(data, custom_fields)
    
    # questionne l'API pour récupérer les données de l'employé 
    collaborateur_info = resquest_collaborateur_info(data.get("Used_by"))
    collaborateurs = [requester for requester in collaborateur_info['requesters'] if requester['active']]
    direction = collaborateurs[0].get('address')
    logging.debug(f'direction : {direction}')
    nom_de_service = collaborateurs[0].get('department_names')[0] if collaborateurs[0].get('department_names') else 'n/c'
    logging.debug(f'nom de service : {nom_de_service}')
    
    # gerer les champs spécifiques si l'actif est un Mobile Phone
    if cl_type == 'Mobile Phone' :
        # Initialise les valeurs reçues dans un dict pour les actifs de type Mobile Phone
        dic = {
            '{{Numéro_de_téléphone}}': '0' + str(
                custom_fields.get('numro_de_tlphone_50000227396', '.....................') or '.....................'),
            '{{IMEI}}': (
                custom_fields.get('imei_number_50000227396', '.....................') or '.....................'),
            '{{PIN}}': (
                custom_fields.get('pin_code_50000227396', '.....................') or '.....................'),
            '{{PUK}}': (
                custom_fields.get('puk_code_50000227396', '.....................') or '.....................'),
            '{{Lock}}': (
                custom_fields.get('lock_code_50000227396', '.....................') or '.....................')
        }
    elif cl_type in ['Tablet', 'Laptop'] :
        # Initialise les valeurs reçues dans un dict pour les actifs de type Tablet ou Laptop
        dic = {
            "{{Modèle}}" : resquest_product_name(custom_fields.get('product_50000227369')),
            "{{Direction}}" : direction,
            "{{Nom du service}}" : nom_de_service,
            "{{Utilise_par1}}" : data.get("Used_by"),
            "{{Nom_T1}}" : data.get("Nom_T"),
            "{{Nom_T}}" : data.get("Nom_T"),
            "{{b_Wupdate}}" : '✓' if data.get("b_Wupdate", False) else '☐',
            "{{b_Lvantage}}" : '✓' if data.get("b_Lvantage", False) else '☐',
            "{{b_Dell}}" : '✓' if data.get("b_Dell", False) else '☐',
            "{{b_Intel}}" : '✓' if data.get("b_Intel", False) else '☐',
            "{{b_Ms}}" : '✓' if data.get("b_Ms", False) else '☐',
            "{{Domaine}}" : data.get("Domaine", "..........") or "..........",
            "{{b_Sophos}}" : '✓' if data.get("b_Sophos", False) else '☐',
            "{{b_Ninite}}" : '✓' if data.get("b_Ninite", False) else '☐',
            "{{b_Of}}" : '✓' if data.get("b_Of", False) else '☐',
            "{{b_Fclient}}" : '✓' if data.get("b_Fclient", False) else '☐',
            "{{b_Edge}}" : '✓' if data.get("b_Edge", False) else '☐',
            "{{b_Teams}}" : '✓' if data.get("b_Teams", False) else '☐',
            "{{b_PRT}}" : '✓' if data.get("b_PRT", False) else '☐',
            "{{b_Poutlook}}" : '✓' if data.get("b_Poutlook", False) else '☐',
            "{{commentaire}}" : ('Commentaire : ' + data.get("commentaire", "") or "") if data.get("commentaire", "") else '',
            "{{Matériel_Sup_list}}" : ('Matériel supplémentaire : ' + data.get("Matériel_Sup_list")) if data.get("Matériel_Sup_list", "") else ""
            }
    else:
        content_note = (f"Le type d'actif reçu n'est pas géré par ce script. Type actif reçu = {cl_type}")
        add_note_to_ticket(TICKET_ID, content_note)
        logging.error(f"script.py :\ncreate_placeholders : {note_content}")
        ERRORS_OCCURED = True
        raise ValueError()
    
    # complète le dict placeholders avec les valeurs et clés (format placeholders) des données envoyé par Freshdeks
    for key, value in dic.items():
            placeholders[key] = value
            
    return placeholders

def replace_placeholders(raw_data):
    """ Remplace les données dans le document template, modifie son nom, et l'enregistre au bon endroit """
    logging.debug("""
    *********************************************************************************************************
               ********************************** Start script.py **********************************
    *********************************************************************************************************""")
    ERRORS_OCCURED = False
    # initialisation du TICKET_ID
    data = create_json(raw_data)
    TICKET_ID = data.get('ticket_id').split('-')[1]
    try:
        # création des JSON
        data = create_json(raw_data)
        custom_fields = create_json(data.get("custom_fields"))
        
        # mise à jour de TICKET_ID
        TICKET_ID = data.get('ticket_id').split('-')[1]
        
        # récupère le type d'actif
        cl_type = data.get('Cl_type')
        
        # Choisi le template en fonction du type d'actif
        doc_path = TEMPLATE[cl_type]

        # créer les champs à insérer dans le doc
        placeholders = create_placeholders(cl_type, data, custom_fields)
        logging.debug(f"Afficher palceholders : {placeholders}")

        # ouvrir doc template pour pouvoir le modifier
        doc = open_doc(doc_path)
        
        # créer les dossiers de destination du fichier rempli
        create_dir('./documents_finaux')
        create_dir(os.path.join('./documents_finaux', placeholders["{{Utilise_par}}"]))
        dossier_nom = os.path.join('./documents_finaux', placeholders["{{Utilise_par}}"])
        
        # Définie les chemins de sauvegarde unique pour les documents .docx et .pdf finaux
        docx_path = get_unique_filename(os.path.join(dossier_nom, f'Attribuer_{cl_type}_{placeholders["{{Asset_tag}}"]}_{placeholders["{{Utilise_par}}"]}.docx'))
        pdf_path = get_unique_filename(os.path.join(dossier_nom, f'Attribuer_{cl_type}_{placeholders["{{Asset_tag}}"]}_{placeholders["{{Utilise_par}}"]}.pdf'))

        # Insérer les valeurs dans le document
        for paragraph in doc.paragraphs:
            for placeholder, value in placeholders.items():
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, str(value))
                    logging.debug(f"remplacement :\t {placeholder} ................ {str(value)}")
        
        # Sauvegarde l'isntance du document modifié
        doc.save(docx_path)

        # Convertir en .PDF
        result = subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', docx_path], capture_output=True, text=True)
        if result.returncode != 0:
            content_note = f"Erreur lors de la conversion en PDF : {result.stderr}"
            logging.error(f"script.py :\nreplace_placeholders : {content_note}")
            add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
            ERRORS_OCCURED = True
            raise ValueError() 
        
        # Récupère le nom du fichier et ajoute .pdf
        pdf_generated_filename = os.path.splitext(os.path.basename(docx_path))[0] + '.pdf'
        # Récupère le chemin vers le fichier .PDF généré
        pdf_generated_path = os.path.join(os.getcwd(), pdf_generated_filename)
        
        # Déplacer le .PDF à pdf_path
        if os.path.exists(pdf_generated_path):
            os.rename(pdf_generated_path, pdf_path)
            if TICKET_ID :
                add_attachment_to_ticket(TICKET_ID, pdf_path)
        else:
            content_note = (f"Le fichier PDF généré n'existe pas :\n {pdf_generated_path}")
            logging.error(f"script.py :\nreplace_placeholders : {note_content}")
            add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, vérifier que l'inventaire soit bien à jour et contsacter l'administrateur si le problème persiste")
            ERRORS_OCCURED = True
            raise FileNotFoundError()
            
        # Effacer le .docx
        if os.path.isfile(docx_path):
            os.remove(docx_path)
        return ERRORS_OCCURED

    except Exception as e:
        content_note = f"script.py :\n Erreur dans replace_placeholders : {e}"
        add_note_to_ticket(TICKET_ID, "Une erreur s'est produite, veuillez contacter votre administrateur si l'erreur se répète")
        logging.error(f"script.py :\nreplace_placeholders : {content_note}", exc_info=True)
        ERRORS_OCCURED = True
        raise ValueError()
        
if __name__ == '__main__':
    logging.info("************** MAIN START ****************")
